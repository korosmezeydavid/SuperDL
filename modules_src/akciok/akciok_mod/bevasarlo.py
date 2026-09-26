# -*- coding: utf-8 -*-
"""Saját bevásárlólista a gépen (offline), és a szinkron a telefonnal.

A GÉPEN: ~/.superdl/bevasarlolista.json – ugyanaz a szerkezet, mint a
telefonon (ShoppingListStore): névvel azonosított listák, a tétel: id, név,
megvan-e, ár forintban. Így a két oldal egy az egyben megfeleltethető.

A TELEFONNAL: az Átjáró portálján át (a telefon helyi WiFi-portálja, 4 jegyű
forgó PIN). Androidon ehhez NEM kellett semmit változtatni – a portál már
tudja: /shopping (lista megnézése), /shopping/select, /shopping/newlist,
/shopping/add, /shopping/toggle. A telefon címét az Átjáró modul beállításából
vesszük, ha van; a PIN minden alkalommal a felhasználótól jön (forog).

Összefésülés: név szerint (kisbetű, ékezet nélkül). Semmit nem törlünk
automatikusan egyik oldalon sem – a szinkron csak ÖSSZEADJA a két listát, és
ha valamelyik oldalon „megvan” a tétel, a másikon is az lesz.
"""
import html
import json
import re
from pathlib import Path

from .termek import Termek, ekezet_nelkul

ALAP_LISTA = "Akciós bevásárlás"
FAJL = Path.home() / ".superdl" / "bevasarlolista.json"
ATJARO_FAJL = Path.home() / ".superdl" / "atjaro.json"
MAX_TETEL = 100                      # a telefon korlátja (MAX_ITEMS)


# ---- helyi lista -----------------------------------------------------------

def betolt() -> dict:
    try:
        d = json.loads(FAJL.read_text(encoding="utf-8"))
        if isinstance(d, dict) and isinstance(d.get("listak"), dict):
            return d
    except (OSError, ValueError):
        pass
    return {"aktiv": ALAP_LISTA, "listak": {ALAP_LISTA: []}}


def ment(adat: dict) -> None:
    FAJL.parent.mkdir(parents=True, exist_ok=True)
    tmp = FAJL.with_suffix(".tmp")
    tmp.write_text(json.dumps(adat, ensure_ascii=False, indent=1),
                   encoding="utf-8")
    tmp.replace(FAJL)


def aktiv_nev(adat: dict) -> str:
    nev = adat.get("aktiv") or ALAP_LISTA
    adat.setdefault("listak", {}).setdefault(nev, [])
    return nev


def tetelek(adat: dict, lista: str | None = None) -> list:
    return adat.get("listak", {}).get(lista or aktiv_nev(adat), [])


def _kulcs(nev: str) -> str:
    return re.sub(r"\s+", " ", ekezet_nelkul(nev)).strip()


def tetel_nev(t: Termek) -> str:
    """A listára kerülő név: a bolt is benne van, mert a boltban az a kérdés,
    HOL van akcióban („Rántott petrella (Penny)”)."""
    return "%s (%s)" % (t.nev, t.bolt)


def hozzaad(adat: dict, nev: str, ar: int | None = None,
            lista: str | None = None) -> tuple:
    """(tétel, új-e). Ugyanaz a név kétszer nem kerül fel."""
    lst = adat.setdefault("listak", {}).setdefault(lista or aktiv_nev(adat), [])
    k = _kulcs(nev)
    for t in lst:
        if _kulcs(t.get("name", "")) == k:
            return t, False
    if len(lst) >= MAX_TETEL:
        raise ValueError("A lista megtelt (legfeljebb %d tétel – ennyit "
                         "enged a telefon is)." % MAX_TETEL)
    uj = {"id": max([t.get("id", 0) for t in lst] + [0]) + 1,
          "name": nev.strip(), "checked": False}
    if ar is not None:
        uj["priceHuf"] = int(ar)
    lst.append(uj)
    return uj, True


def termek_hozzaad(adat: dict, t: Termek) -> tuple:
    return hozzaad(adat, tetel_nev(t), t.legjobb_ar())


def torol(adat: dict, tetel_id: int, lista: str | None = None) -> bool:
    lst = tetelek(adat, lista)
    for i, t in enumerate(lst):
        if t.get("id") == tetel_id:
            del lst[i]
            return True
    return False


def megvan_valt(adat: dict, tetel_id: int, lista: str | None = None):
    for t in tetelek(adat, lista):
        if t.get("id") == tetel_id:
            t["checked"] = not t.get("checked", False)
            return t
    return None


def osszeg(lst: list) -> int:
    return sum(int(t.get("priceHuf") or 0) for t in lst)


_HONAPOK = ("január", "február", "március", "április", "május", "június",
            "július", "augusztus", "szeptember", "október", "november",
            "december")


def _ft_szep(n: int) -> str:
    """5480 → „5 480" (az ezres tagolás a papíron és felolvasva is jobb)."""
    return "{:,}".format(int(n)).replace(",", " ")


def szoveges(adat: dict, ma=None) -> str:
    """A lista KIKÜLDHETŐ szövege (vágólap, fájl, Super Edit) – Petrus
    József kérése: „el tudom küldeni messengerben vagy e-mailben… a
    segítőnek, aki bevásárol nekem".

    Soronként egy tétel a bolttal és az árral; ami már MEGVAN, az külön, a
    végén – a segítőnek csak az marad a listán, amit még meg kell venni."""
    import datetime as _dt
    ma = ma or _dt.date.today()
    lst = tetelek(adat)
    kell = [t for t in lst if not t.get("checked")]
    megvan = [t for t in lst if t.get("checked")]
    sorok = ["Bevásárlólista – %s (%d. %s %d.)" % (
        aktiv_nev(adat), ma.year, _HONAPOK[ma.month - 1], ma.day), ""]
    if not kell:
        sorok.append("Minden megvan a listáról.")
    for t in kell:
        ar = t.get("priceHuf")
        sorok.append("%s%s" % (t.get("name", ""),
                               " – %s Ft" % _ft_szep(ar) if ar else ""))
    if kell and any(t.get("priceHuf") for t in kell):
        sorok += ["", "Összesen kb. %s Ft" % _ft_szep(osszeg(kell))]
    if megvan:
        sorok += ["", "Már megvan: " + ", ".join(t.get("name", "") for t in megvan)]
    return "\n".join(sorok) + "\n"


def ment_fajlba(szoveg: str, ut: str) -> None:
    """.docx → Word-dokumentum (az első sor címsorként); minden más → UTF-8
    szövegfájl Windows-sorvégekkel (a Jegyzettömb és a Super Edit is jól
    olvassa)."""
    if ut.lower().endswith(".docx"):
        import docx                       # a Core-ban (python-docx)
        doc = docx.Document()
        sorok = szoveg.rstrip("\n").split("\n")
        doc.add_heading(sorok[0], level=1)
        for s in sorok[1:]:
            doc.add_paragraph(s)
        doc.save(ut)
        return
    with open(ut, "w", encoding="utf-8-sig", newline="\r\n") as f:
        f.write(szoveg)


def tetel_sor(t: dict) -> str:
    ar = t.get("priceHuf")
    return "%s%s. %s" % (t.get("name", ""),
                         ", %d forint" % ar if ar is not None else "",
                         "Megvan" if t.get("checked") else "Még nincs meg")


# ---- telefon (Átjáró-portál) ---------------------------------------------

def telefon_cim() -> tuple:
    """(ip, port) az Átjáró beállításából, vagy ("", 8080)."""
    try:
        d = json.loads(ATJARO_FAJL.read_text(encoding="utf-8"))
        return (d.get("ip") or "").strip(), int(d.get("port") or 8080)
    except (OSError, ValueError, TypeError):
        return "", 8080


def telefon_cim_ment(ip: str, port: int = 8080) -> None:
    try:
        d = json.loads(ATJARO_FAJL.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        d = {}
    d["ip"], d["port"] = (ip or "").strip(), int(port or 8080)
    ATJARO_FAJL.parent.mkdir(parents=True, exist_ok=True)
    ATJARO_FAJL.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")


class RosszPin(RuntimeError):
    pass


def oldal_elemez(szoveg: str) -> dict:
    """A telefon /shopping oldala → {"listak": [...], "aktiv": név|None,
    "tetelek": [{"id","name","checked","priceHuf"}]}."""
    if 'name="pin"' in szoveg and "shop_newlist" not in szoveg:
        raise RosszPin("Rossz PIN – a telefon a bejelentkezést kérte.")
    if "shop_newlist" not in szoveg:
        raise RuntimeError("A telefon nem a bevásárlólistát küldte vissza – "
                           "lehet, hogy régebbi a SuperDL a telefonon.")
    listak, aktiv = [], None
    for m in re.finditer(r'<option value="([^"]*)"([^>]*)>', szoveg):
        n = html.unescape(m.group(1))
        listak.append(n)
        if "selected" in m.group(2):
            aktiv = n
    tetelek = []
    for li in re.findall(r'<li class="media-item">(.*?)</li>', szoveg, re.S):
        nev = re.search(r'<span class="media-name">(.*?)</span>', li, re.S)
        meret = re.search(r'<span class="media-size">(.*?)</span>', li, re.S)
        azon = re.search(r'name="id" value="(\d+)"', li)
        if not (nev and azon):
            continue
        nyers = nev.group(1)
        megvan = "<s>" in nyers
        ar = None
        if meret:
            m = re.search(r"(\d+)\s*Ft", meret.group(1))
            ar = int(m.group(1)) if m else None
        tetelek.append({"id": int(azon.group(1)),
                        "name": html.unescape(re.sub(r"<[^>]+>", "", nyers)).strip(),
                        "checked": megvan, "priceHuf": ar})
    if aktiv is None and len(listak) == 1:
        aktiv = listak[0]
    return {"listak": listak, "aktiv": aktiv, "tetelek": tetelek}


class Telefon:
    """A telefon bevásárlólistája a portálon át. `requests` a Core-ból."""

    def __init__(self, ip: str, pin: str, port: int = 8080, session=None):
        ip = (ip or "").strip()
        self.alap = ip.rstrip("/") if ip.startswith("http") \
            else "http://%s:%d" % (ip, int(port or 8080))
        self.pin = str(pin or "").strip()
        if session is None:
            import requests
            session = requests.Session()
        self.s = session

    def _get(self):
        r = self.s.get(self.alap + "/shopping", params={"pin": self.pin},
                       timeout=15)
        r.raise_for_status()
        return oldal_elemez(r.text)

    def _post(self, ut, adat):
        r = self.s.post(self.alap + ut, params={"pin": self.pin}, data=adat,
                        timeout=15)
        r.raise_for_status()
        return oldal_elemez(r.text)

    def lista_megnyit(self, nev: str) -> dict:
        """A telefonon a megadott listát teszi aktívvá (ha nincs, létrehozza)."""
        allapot = self._get()
        if nev in allapot["listak"]:
            if allapot["aktiv"] != nev:
                allapot = self._post("/shopping/select", {"list": nev})
        else:
            allapot = self._post("/shopping/newlist", {"name": nev})
        if allapot["aktiv"] != nev:
            raise RuntimeError("A telefonon nem sikerült megnyitni a(z) „%s” "
                               "listát." % nev)
        return allapot

    def szinkron(self, adat: dict, lista: str | None = None) -> dict:
        """Kétirányú összefésülés. Vissza: {"le": db, "fel": db, "pipa": db}."""
        lista = lista or aktiv_nev(adat)
        allapot = self.lista_megnyit(lista)
        telefonon = {_kulcs(t["name"]): t for t in allapot["tetelek"]}
        helyi = tetelek(adat, lista)
        eredmeny = {"le": 0, "fel": 0, "pipa": 0}
        # 1) ami csak a telefonon van → a gépre
        for k, t in telefonon.items():
            if not any(_kulcs(x.get("name", "")) == k for x in helyi):
                uj, _ = hozzaad(adat, t["name"], t.get("priceHuf"), lista)
                uj["checked"] = bool(t.get("checked"))
                eredmeny["le"] += 1
        # 2) ami csak a gépen van → a telefonra
        for x in list(tetelek(adat, lista)):
            k = _kulcs(x.get("name", ""))
            if k in telefonon:
                continue
            adat_ = {"name": x["name"]}
            if x.get("priceHuf") is not None:
                adat_["price"] = str(x["priceHuf"])
            allapot = self._post("/shopping/add", adat_)
            eredmeny["fel"] += 1
            uj = next((t for t in allapot["tetelek"]
                       if _kulcs(t["name"]) == k), None)
            if uj and x.get("checked"):
                self._post("/shopping/toggle", {"id": str(uj["id"])})
            if uj:
                telefonon[k] = dict(uj, checked=bool(x.get("checked")))
        # 3) „megvan”: ha bármelyik oldalon megvan, mindkettőn az legyen
        for x in tetelek(adat, lista):
            t = telefonon.get(_kulcs(x.get("name", "")))
            if not t:
                continue
            if t.get("checked") and not x.get("checked"):
                x["checked"] = True
                eredmeny["pipa"] += 1
            elif x.get("checked") and not t.get("checked"):
                self._post("/shopping/toggle", {"id": str(t["id"])})
                eredmeny["pipa"] += 1
        return eredmeny
