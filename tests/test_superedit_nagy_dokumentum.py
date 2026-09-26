# -*- coding: utf-8 -*-
"""Super Edit 1.0.1 – nagy dokumentum mentése és megnyitása.

Farkas István hibajelentése (2026-09-25): egy Word-mentés ~40 másodpercre
megakasztotta az egész programot. Két javítás: (1) a .docx írása és olvasása
közvetlenül az XML-fán, keresések nélkül – mérve: 40 000 bekezdés mentése
~36 mp → ~2 mp, megnyitása ~90 mp → ~2 mp; (2) a mentés külön szálon fut,
az ablak közben válaszol."""
import os
import re
import time
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
FM = pytest.importorskip("modules_src.superedit.superedit_mod.formatum")

FORRAS = Path("modules_src/superedit/superedit_mod/szerkesztowin.py")


def _nagy(n):
    ki = []
    for i in range(n):
        if i % 50 == 0:
            ki.append(("Fejezet %d" % i, "Címsor 1"))
        else:
            ki.append(("x", "Normál", [("Első rész ", False, False, False),
                                       ("félkövér", True, False, False),
                                       (" vége.", False, True, True)]))
    return ki


def test_nagy_dokumentum_gyorsan_ment_es_nyit(tmp_path):
    ut = str(tmp_path / "nagy.docx")
    bek = _nagy(5000)
    t = time.perf_counter()
    FM.ment(ut, bek)
    mentes = time.perf_counter() - t
    t = time.perf_counter()
    vissza = FM.megnyit(ut).bekezdesek
    nyitas = time.perf_counter() - t
    # a régi út 5000 bekezdésnél ~4,5 mp mentés és ~11 mp nyitás volt
    assert mentes < 2.5 and nyitas < 2.5, (mentes, nyitas)
    assert len(vissza) == 5000
    assert vissza[0][:2] == ("Fejezet 0", "Címsor 1")
    assert vissza[1][2] == [("Első rész ", False, False, False),
                            ("félkövér", True, False, False),
                            (" vége.", False, True, True)]


def test_word_ugyanazt_latja(tmp_path):
    """A kész fájlt a python-docx saját útján is visszaolvassuk: a Word
    szempontjából ugyanaz a dokumentum, mint amit a régi mentés írt."""
    ut = str(tmp_path / "a.docx")
    FM.ment(ut, [("Cím", "Címsor 2"),
                 ("a\tb\nc", "Normál", [("a\tb\nc", False, False, False)]),
                 ("", "Normál")])
    d = docx.Document(ut)
    ps = d.paragraphs
    assert ps[0].style.name == "Heading 2" and ps[0].text == "Cím"
    assert ps[1].text == "a\tb\nc"
    assert ps[2].text == ""


def test_tiltott_vezerlokarakter_nem_dontja_el_a_mentest(tmp_path):
    ut = str(tmp_path / "b.docx")
    FM.ment(ut, [("PDF\x0bből\x01", "Normál")])
    assert FM.megnyit(ut).bekezdesek[0][0] == "PDFből"


def test_sikertelen_mentes_nem_hagy_felkesz_fajlt(tmp_path, monkeypatch):
    ut = tmp_path / "c.docx"
    FM.ment(str(ut), [("eredeti", "Normál")])

    def rossz(self, path):
        Path(path).write_bytes(b"felkesz")
        raise OSError("tele a lemez")
    monkeypatch.setattr(docx.document.Document, "save", rossz)
    with pytest.raises(OSError):
        FM.ment(str(ut), [("új", "Normál")])
    assert list(tmp_path.iterdir()) == [ut]          # nincs .tmp maradék
    monkeypatch.undo()
    assert FM.megnyit(str(ut)).bekezdesek[0][0] == "eredeti"


def test_hivatkozas_szovege_is_megmarad(tmp_path):
    ut = str(tmp_path / "h.docx")
    FM.ment(ut, [("Lásd: ", "Normál")])
    d = docx.Document(ut)
    from docx.oxml.ns import qn
    from lxml import etree
    p = d.paragraphs[0]._p
    hl = etree.SubElement(p, qn("w:hyperlink"))
    r = etree.SubElement(hl, qn("w:r"))
    etree.SubElement(r, qn("w:t")).text = "super-dl.com"
    d.save(ut)
    assert FM.megnyit(ut).bekezdesek[0][0] == "Lásd: super-dl.com"


def test_a_mentes_hatterszalon_fut():
    src = FORRAS.read_text(encoding="utf-8")
    assert "def _hatter_ment" in src and "threading.Thread" in src
    for fv in ("_ment", "_ment_maskent", "_pdf_export"):
        test = re.search(r"def %s\(self.*?(?=\n    def )" % fv, src, re.S)
        assert test and "_hatter_ment(" in test.group(0), fv
        assert "FM.ment(" not in test.group(0), fv
    # mentés közben nem zárható be az ablak (félkész fájl lenne)
    assert '_mentes_fut", False) and e.CanVeto()' in src
